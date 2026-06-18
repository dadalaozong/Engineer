import json
from flask import Blueprint, render_template, request, Response, jsonify

bp = Blueprint("ai_writer", __name__, url_prefix="/ai")

_DOC_TYPES = [
    ("work_summary", "个人工作总结"),
    ("masterwork",   "代表作工程业绩"),
    ("achievement",  "科技成果说明"),
]

@bp.route("/")
def index_page():
    return render_template("ai/writer.html", doc_types=_DOC_TYPES, prefill={})

@bp.route("/with-applicant")
def writer_with_applicant():
    """AI写作页面，自动从申报人/项目数据中预填字段。"""
    from database.models import get_applicant, get_project, list_achievements, list_papers
    from datetime import date

    aid = request.args.get("applicant_id", type=int)
    pid = request.args.get("project_id", type=int)
    doc_type = request.args.get("doc_type", "work_summary")

    applicant = get_applicant(aid) if aid else {}
    project   = get_project(pid) if pid else {}
    a = applicant or {}
    p = dict(project) if project else {}

    cur_year = date.today().year
    work_age = ""
    if a.get("work_start_year"):
        try: work_age = str(cur_year - int(a["work_start_year"]))
        except Exception: pass

    achievements = list_achievements(aid) if aid else []
    papers       = list_papers(aid) if aid else []
    proj_names   = "、".join(r["project_name"] for r in achievements[:3] if r.get("project_name"))
    paper_titles = "、".join(r["title"] for r in papers[:2] if r.get("title"))

    prefill = {
        "work_summary": {
            "name":           a.get("name",""),
            "years":          work_age,
            "specialties":    a.get("title_specialty") or a.get("current_specialty",""),
            "major_projects": proj_names,
            "achievements":   paper_titles or a.get("notes",""),
        },
        "masterwork": {
            "name":         a.get("name",""),
            "project_name": achievements[0].get("project_name","") if achievements else "",
            "project_type": achievements[0].get("project_type","") if achievements else "",
            "project_scale":achievements[0].get("scale","") if achievements else "",
            "role":         achievements[0].get("role","") if achievements else "",
            "duration":     "",
        },
        "achievement": {
            "name":             a.get("name",""),
            "achievement_name": paper_titles or "",
            "achievement_type": "论文" if papers else "",
            "description":      "",
        },
    }

    return render_template("ai/writer.html", doc_types=_DOC_TYPES,
                           prefill=prefill, active_doc_type=doc_type,
                           applicant=a, project=p)

@bp.route("/generate", methods=["POST"])
def generate():
    data = request.json or {}
    doc_type = data.get("doc_type", "work_summary")
    fields = data.get("info") or data.get("fields") or {}

    def stream():
        try:
            from core.ai_writer import stream_write
            for chunk in stream_write(doc_type, fields):
                yield f"data: {json.dumps({'text': chunk}, ensure_ascii=False)}\n\n"
        except Exception as e:
            mock = (
                f"个人工作总结\n\n"
                f"本人{fields.get('name','XXX')}，从事专业技术工作{fields.get('years','多')}年。\n\n"
                f"一、主要工作内容\n\n认真履行岗位职责，参与完成多项重要工程项目，"
                f"包括：{fields.get('major_projects','各类项目')}。\n\n"
                f"二、主要业绩成果\n\n{fields.get('achievements','在工作中取得了一定成绩。')}\n\n"
                f"三、下一步计划\n\n继续加强学习，提升专业技术水平，为单位发展贡献力量。"
            )
            for ch in mock:
                yield f"data: {json.dumps({'text': ch}, ensure_ascii=False)}\n\n"
        yield f"data: [DONE]\n\n"

    return Response(stream(), mimetype="text/event-stream",
                    headers={"X-Accel-Buffering": "no", "Cache-Control": "no-cache"})

@bp.route("/export", methods=["POST"])
def export_doc():
    data = request.json or {}
    content = data.get("content", "")
    doc_type = data.get("doc_type", "文档")
    label = next((lbl for key, lbl in _DOC_TYPES if key == doc_type), doc_type)
    try:
        from docx import Document
        from io import BytesIO
        doc = Document()
        doc.add_heading(label, level=1)
        for para in content.split("\n"):
            doc.add_paragraph(para)
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        from flask import make_response
        resp = make_response(buf.read())
        resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        resp.headers["Content-Disposition"] = f'attachment; filename="{label}.docx"'
        return resp
    except ImportError:
        from flask import make_response
        resp = make_response(content.encode("utf-8"))
        resp.headers["Content-Type"] = "text/plain; charset=utf-8"
        resp.headers["Content-Disposition"] = f'attachment; filename="{label}.txt"'
        return resp
