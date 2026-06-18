from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify
from database.models import (
    list_applicants, get_applicant, insert_applicant, update_applicant, delete_applicant,
    list_achievements, get_achievement, insert_achievement, update_achievement, delete_achievement,
    list_awards, insert_award, update_award, delete_award,
    list_papers, insert_paper, update_paper, delete_paper,
    list_contact_logs, insert_contact_log, delete_contact_log,
)

bp = Blueprint("applicants", __name__, url_prefix="/applicants")

_A_COLS = [
    "name","id_card","gender","birth_date","phone","email","ethnicity",
    "education","major","school","graduation_year",
    "work_unit","work_unit_type","work_start_year",
    "title_level","title_year","title_month","notes"
]

def _form(cols):
    return {c: (request.form.get(c) or "").strip() for c in cols}

# ── 列表 / 新增 / 编辑 / 删除 ────────────────────────────────────

@bp.route("/")
def list_page():
    q = request.args.get("q", "")
    applicants = list_applicants(q)
    return render_template("applicants/list.html", applicants=applicants)

@bp.route("/new", methods=["GET"])
def new_page():
    return render_template("applicants/form.html", applicant=None)

@bp.route("/new", methods=["POST"])
def create():
    insert_applicant(**_form(_A_COLS))
    flash("申报人已添加", "success")
    return redirect(url_for("applicants.list_page"))

@bp.route("/new-json", methods=["POST"])
def create_json():
    """AJAX接口 — OCR弹窗保存申报人。"""
    d    = request.get_json(silent=True) or {}
    name = (d.get("name") or "").strip()
    if not name:
        return jsonify({"success": False, "error": "姓名不能为空"})
    fields = {c: (d.get(c) or "") for c in _A_COLS}
    new_id = insert_applicant(**fields)
    return jsonify({"success": True, "id": new_id})

@bp.route("/<int:aid>/edit", methods=["GET"])
def edit_page(aid):
    applicant = get_applicant(aid)
    if not applicant:
        flash("申报人不存在", "danger")
        return redirect(url_for("applicants.list_page"))
    return render_template("applicants/form.html", applicant=applicant)

@bp.route("/<int:aid>/edit", methods=["POST"])
def update(aid):
    update_applicant(aid, **_form(_A_COLS))
    flash("申报人信息已更新", "success")
    return redirect(url_for("applicants.detail", aid=aid))

@bp.route("/<int:aid>/delete", methods=["POST"])
def delete(aid):
    delete_applicant(aid)
    flash("申报人已删除", "success")
    return redirect(url_for("applicants.list_page"))

# ── 详情页（业绩/获奖/论文/沟通记录汇总）───────────────────────

@bp.route("/<int:aid>")
def detail(aid):
    applicant    = get_applicant(aid)
    if not applicant:
        flash("申报人不存在", "danger")
        return redirect(url_for("applicants.list_page"))
    achievements = list_achievements(aid)
    awards       = list_awards(aid)
    papers       = list_papers(aid)
    logs         = list_contact_logs(aid)
    from datetime import date
    return render_template("applicants/detail.html",
                           applicant=applicant,
                           achievements=achievements,
                           awards=awards,
                           papers=papers,
                           logs=logs,
                           today=date.today().isoformat())

# ── 工程业绩 CRUD ─────────────────────────────────────────────────

_ACH_COLS = [
    "project_name","project_type","scale","role","start_date","end_date",
    "location","owner","contractor","investment","area","description","is_representative"
]

@bp.route("/<int:aid>/achievements/new", methods=["POST"])
def achievement_create(aid):
    d = _form(_ACH_COLS)
    d["applicant_id"] = aid
    d["is_representative"] = 1 if request.form.get("is_representative") else 0
    insert_achievement(**d)
    flash("工程业绩已添加", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#achievements")

@bp.route("/<int:aid>/achievements/<int:rid>/edit", methods=["POST"])
def achievement_update(aid, rid):
    d = _form(_ACH_COLS)
    d["is_representative"] = 1 if request.form.get("is_representative") else 0
    update_achievement(rid, **d)
    flash("已保存", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#achievements")

@bp.route("/<int:aid>/achievements/<int:rid>/delete", methods=["POST"])
def achievement_delete(aid, rid):
    delete_achievement(rid)
    flash("已删除", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#achievements")

# ── 获奖 CRUD ─────────────────────────────────────────────────────

_AWD_COLS = ["award_name","award_level","award_year","award_org","rank","certificate_no","notes"]

@bp.route("/<int:aid>/awards/new", methods=["POST"])
def award_create(aid):
    d = _form(_AWD_COLS); d["applicant_id"] = aid
    insert_award(**d)
    flash("获奖信息已添加", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#awards")

@bp.route("/<int:aid>/awards/<int:wid>/edit", methods=["POST"])
def award_update(aid, wid):
    update_award(wid, **_form(_AWD_COLS))
    flash("已保存", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#awards")

@bp.route("/<int:aid>/awards/<int:wid>/delete", methods=["POST"])
def award_delete(aid, wid):
    delete_award(wid)
    flash("已删除", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#awards")

# ── 论文 CRUD ─────────────────────────────────────────────────────

_PAP_COLS = ["title","paper_type","journal","pub_date","author_rank","cn_issn","notes"]

@bp.route("/<int:aid>/papers/new", methods=["POST"])
def paper_create(aid):
    d = _form(_PAP_COLS); d["applicant_id"] = aid
    insert_paper(**d)
    flash("论文/著作已添加", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#papers")

@bp.route("/<int:aid>/papers/<int:pid>/edit", methods=["POST"])
def paper_update(aid, pid):
    update_paper(pid, **_form(_PAP_COLS))
    flash("已保存", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#papers")

@bp.route("/<int:aid>/papers/<int:pid>/delete", methods=["POST"])
def paper_delete(aid, pid):
    delete_paper(pid)
    flash("已删除", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#papers")

# ── 沟通记录 ──────────────────────────────────────────────────────

_LOG_COLS = ["contact_date","contact_type","content","follow_up","follow_up_date","operator"]

@bp.route("/<int:aid>/logs/new", methods=["POST"])
def log_create(aid):
    d = _form(_LOG_COLS); d["applicant_id"] = aid
    insert_contact_log(**d)
    flash("沟通记录已添加", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#logs")

@bp.route("/<int:aid>/logs/<int:lid>/delete", methods=["POST"])
def log_delete(aid, lid):
    delete_contact_log(lid)
    flash("已删除", "success")
    return redirect(url_for("applicants.detail", aid=aid) + "#logs")

# ── Excel 导出 ─────────────────────────────────────────────────────

@bp.route("/export/excel")
def export_excel():
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from flask import make_response

    q = request.args.get("q", "")
    rows = list_applicants(q)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "申报人列表"

    headers = ["姓名","身份证号","性别","出生日期","电话","邮箱","工作单位",
               "学历","专业","毕业院校","毕业年份","参加工作年份",
               "现职称等级","取得年月","备注"]
    keys    = ["name","id_card","gender","birth_date","phone","email","work_unit",
               "education","major","school","graduation_year","work_start_year",
               "title_level","title_month","notes"]

    hdr_fill = PatternFill("solid", fgColor="2563EB")
    hdr_font = Font(color="FFFFFF", bold=True, size=11)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(12, len(h) * 2 + 2)

    for r, row in enumerate(rows, 2):
        for col, key in enumerate(keys, 1):
            ws.cell(row=r, column=col, value=row.get(key) or "")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = make_response(buf.read())
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = "attachment; filename=applicants.xlsx"
    return resp
