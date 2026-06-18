from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify
from database.models import (
    list_projects, get_project, insert_project, update_project, delete_project,
    list_applicants, list_batches, update_project_stage,
    get_applicant, list_achievements, list_awards, list_papers,
    list_social_insurance, list_edu_trainings,
)

bp = Blueprint("projects", __name__, url_prefix="/projects")

SPECIALTY_NAMES = ["房屋建筑工程", "市政公用工程", "装饰装修工程", "机电安装工程", "公路工程"]

def _form_data():
    return dict(
        applicant_id=request.form.get("applicant_id", ""),
        batch_id=request.form.get("batch_id", ""),
        industry=request.form.get("industry", ""),
        committee=request.form.get("committee", ""),
        apply_level=request.form.get("apply_level", ""),
        specialty=request.form.get("specialty", ""),
        project_name=request.form.get("project_name", ""),
        project_type=request.form.get("project_type", ""),
        scale=request.form.get("scale", ""),
        role=request.form.get("role", ""),
        folder_path=request.form.get("folder_path", ""),
        status=request.form.get("status", ""),
        notes=request.form.get("notes", ""),
    )

@bp.route("/")
def list_page():
    applicant_id = request.args.get("applicant_id", type=int)
    rows = list_projects(applicant_id=applicant_id)
    return render_template("projects/list.html",
                           rows=rows,
                           applicants=list_applicants(),
                           filter_aid=applicant_id)

@bp.route("/new", methods=["GET"])
def new_page():
    return render_template("projects/form.html", project=None, applicants=list_applicants(), batches=list_batches())

@bp.route("/new", methods=["POST"])
def create():
    insert_project(**_form_data())
    flash("项目已添加", "success")
    return redirect(url_for("projects.list_page"))

@bp.route("/<int:pid>/edit", methods=["GET"])
def edit_page(pid):
    project = get_project(pid)
    return render_template("projects/form.html", project=project, applicants=list_applicants(), batches=list_batches())

@bp.route("/<int:pid>/edit", methods=["POST"])
def update(pid):
    update_project(pid, **_form_data())
    flash("项目已更新", "success")
    return redirect(url_for("projects.list_page"))

@bp.route("/<int:pid>/delete", methods=["POST"])
def delete(pid):
    delete_project(pid)
    flash("项目已删除", "success")
    return redirect(url_for("projects.list_page"))

@bp.route("/scale", methods=["GET"])
def scale_page():
    return render_template("projects/scale.html", SPECIALTY_NAMES=SPECIALTY_NAMES)

def judge(specialty, indicators):
    cost = float(indicators.get("cost") or 0)
    area = float(indicators.get("area") or 0)
    height = float(indicators.get("height") or 0)
    floors = float(indicators.get("floors") or 0)
    road_len = float(indicators.get("road_len") or 0)
    pipe_dia = float(indicators.get("pipe_dia") or 0)
    bridge_span = float(indicators.get("bridge_span") or 0)
    deco_area = float(indicators.get("deco_area") or 0)
    transformer = float(indicators.get("transformer") or 0)
    chiller = float(indicators.get("chiller") or 0)
    speed = float(indicators.get("speed") or 0)

    scale = "小型"
    color = "#6b7280"
    detail = ""

    if specialty == "房屋建筑工程":
        if cost >= 10000 or area >= 10 or height >= 100 or floors >= 3:
            scale, color = "大型", "#dc2626"
            detail = f"造价{cost}万元/面积{area}万㎡/高度{height}m/地下{floors}层"
        elif cost >= 3000 or area >= 3 or height >= 50 or floors >= 2:
            scale, color = "中型", "#d97706"
            detail = f"造价{cost}万元/面积{area}万㎡/高度{height}m/地下{floors}层"
        else:
            detail = f"造价{cost}万元/面积{area}万㎡"
    elif specialty == "市政公用工程":
        if cost >= 5000 or road_len >= 10 or pipe_dia >= 1000 or bridge_span >= 150:
            scale, color = "大型", "#dc2626"
        elif cost >= 1000 or road_len >= 3 or pipe_dia >= 500 or bridge_span >= 50:
            scale, color = "中型", "#d97706"
        detail = f"造价{cost}万元/道路{road_len}km"
    elif specialty == "装饰装修工程":
        if cost >= 2000 or deco_area >= 20000:
            scale, color = "大型", "#dc2626"
        elif cost >= 500 or deco_area >= 5000:
            scale, color = "中型", "#d97706"
        detail = f"造价{cost}万元/面积{deco_area}㎡"
    elif specialty == "机电安装工程":
        if cost >= 3000 or transformer >= 10000 or chiller >= 1000:
            scale, color = "大型", "#dc2626"
        elif cost >= 1000 or transformer >= 3000 or chiller >= 300:
            scale, color = "中型", "#d97706"
        detail = f"造价{cost}万元/变压器{transformer}kVA"
    elif specialty == "公路工程":
        if cost >= 5000 or road_len >= 20 or speed >= 120:
            scale, color = "大型", "#dc2626"
        elif cost >= 1000 or road_len >= 5 or speed >= 80:
            scale, color = "中型", "#d97706"
        detail = f"造价{cost}万元/长度{road_len}km/速度{speed}km/h"

    return scale, color, detail

@bp.route("/scale", methods=["POST"])
def scale_judge():
    data = request.json or {}
    specialty = data.get("specialty", "")
    indicators = data.get("indicators", {})
    scale, color, detail = judge(specialty, indicators)
    return jsonify({"scale": scale, "color": color, "detail": detail})

@bp.route("/<int:pid>/stage", methods=["POST"])
def set_stage(pid):
    d = request.get_json(silent=True) or {}
    stage = int(d.get("stage", 0))
    update_project_stage(pid, stage)
    return jsonify({"ok": True, "stage": stage})

@bp.route("/checker")
def checker_page():
    projects = list_projects()
    return render_template("projects/checker.html", projects=projects)

# ── 材料清单（按广西职称网左侧菜单Tab结构）────────────────────────
#
# fill: "auto"  = 系统从数据库自动填写，无需提供扫描件
#       "file"  = 需申报人提供扫描件放入对应目录
#       "ai"    = 由AI写作模块生成
#
# folder: 对应 core/folder_manager.py SUBFOLDERS 中的目录名
#
# required_for: 哪些申报级别是必填项（空列表=所有级别）
#
_TAB_MATERIALS = [
    {
        "tab": "1",
        "tab_name": "基本信息",
        "folder": "01_基本信息",
        "items": [
            {"name": "身份证国徽面扫描件",        "fill": "file", "required_for": [],
             "note": "OCR识别证件号、有效期、签发机关"},
            {"name": "身份证人像面扫描件",        "fill": "file", "required_for": [],
             "note": "OCR识别姓名、性别、民族、出生日期、地址"},
            {"name": "近期免冠照片",              "fill": "file", "required_for": [],
             "note": "证件照，白底或蓝底，JPG格式，网站左上角上传"},
            {"name": "评审年度",                  "fill": "auto", "required_for": [],
             "note": "从项目年度自动填写"},
            {"name": "参加评审会",                "fill": "auto", "required_for": [],
             "note": "从项目评委会字段自动填写"},
            {"name": "曾用名",                    "fill": "auto", "required_for": [],
             "note": "从申报人档案填写，如无可留空"},
            {"name": "填表时间",                  "fill": "auto", "required_for": [],
             "note": "自动取当前年月"},
            {"name": "证件号码",                  "fill": "auto", "required_for": [],
             "note": "OCR识别后自动填写"},
            {"name": "个人身份性质",              "fill": "auto", "required_for": [],
             "note": "从申报人档案填写（非公单位/公单位等）"},
            {"name": "性别",                      "fill": "auto", "required_for": [],
             "note": "OCR识别后自动填写"},
            {"name": "民族",                      "fill": "auto", "required_for": [],
             "note": "OCR识别后自动填写"},
            {"name": "出生年月",                  "fill": "auto", "required_for": [],
             "note": "OCR识别后自动填写"},
            {"name": "籍贯",                      "fill": "auto", "required_for": [],
             "note": "从申报人档案填写"},
            {"name": "政治面貌",                  "fill": "auto", "required_for": [],
             "note": "从申报人档案填写"},
            {"name": "参加工作时间",              "fill": "auto", "required_for": [],
             "note": "从申报人档案参加工作年月填写"},
            {"name": "联系电话",                  "fill": "auto", "required_for": [],
             "note": "从申报人档案手机号填写"},
            {"name": "电子邮箱",                  "fill": "auto", "required_for": [],
             "note": "从申报人档案邮箱填写"},
            {"name": "拟评职称系列",              "fill": "auto", "required_for": [],
             "note": "从项目申报系列自动填写"},
            {"name": "拟评级别",                  "fill": "auto", "required_for": [],
             "note": "从项目申报级别自动填写"},
            {"name": "拟评专业技术资格",          "fill": "auto", "required_for": [],
             "note": "从项目拟评资格名称自动填写"},
            {"name": "学科",                      "fill": "auto", "required_for": [],
             "note": "从项目学科字段自动填写"},
            {"name": "专业技术工作年限",          "fill": "auto", "required_for": [],
             "note": "系统根据参加工作年月自动计算"},
            {"name": "拟评专业",                  "fill": "auto", "required_for": [],
             "note": "从项目专业字段自动填写"},
            {"name": "申报方式",                  "fill": "auto", "required_for": [],
             "note": "从项目申报方式填写（晋升/初定等）"},
            {"name": "是否第一次申报",            "fill": "auto", "required_for": [],
             "note": "从申报人档案填写"},
            {"name": "曾申报次数",                "fill": "auto", "required_for": [],
             "note": "从申报人档案填写"},
            {"name": "上一次申报时间",            "fill": "auto", "required_for": [],
             "note": "从申报人档案填写，如无可留空"},
            {"name": "参加乡村振兴定向评价",      "fill": "auto", "required_for": [],
             "note": "从申报人档案填写（是/否）"},
            {"name": "是否高技能人才",            "fill": "auto", "required_for": [],
             "note": "从申报人档案填写（是/否）"},
            {"name": "从事技术技能工作年限",      "fill": "auto", "required_for": [],
             "note": "从申报人档案填写，如不适用可留空"},
            {"name": "单位级别",                  "fill": "auto", "required_for": [],
             "note": "从申报人档案工作单位级别填写（市属/省属等）"},
            {"name": "行政职务",                  "fill": "auto", "required_for": [],
             "note": "从申报人档案行政职务填写"},
            {"name": "行政职务任命时间",          "fill": "auto", "required_for": [],
             "note": "从申报人档案行政职务任命时间填写"},
            {"name": "行政职务说明",              "fill": "auto", "required_for": [],
             "note": "从申报人档案行政职务说明填写，如无可留空"},
            {"name": "档案所在地机构名称",        "fill": "auto", "required_for": [],
             "note": "从申报人档案填写"},
            {"name": "联系地址",                  "fill": "auto", "required_for": [],
             "note": "从申报人档案地址填写"},
        ],
    },
    {
        "tab": "2",
        "tab_name": "学历情况",
        "folder": "02_学历情况",
        "items": [
            # ── 扫描件（放入02_学历情况目录）──
            {"name": "最高学历证书扫描件",        "fill": "file", "required_for": [],
             "note": "OCR识别毕业时间/学历/毕业学校/专业"},
            {"name": "学位证书扫描件",            "fill": "file", "required_for": [],
             "note": "有学位的需提供（学士/硕士/博士）"},
            # ── 系统自动填写字段（对应网站列表各列）──
            {"name": "数据来源",                  "fill": "auto", "required_for": [],
             "note": "固定填写【自填】"},
            {"name": "毕业时间",                  "fill": "auto", "required_for": [],
             "note": "格式 YYYY-MM，从档案毕业年月自动填写"},
            {"name": "学历项目",                  "fill": "auto", "required_for": [],
             "note": "最高学历/第二学历等，默认填【最高学历】"},
            {"name": "学历",                      "fill": "auto", "required_for": [],
             "note": "大专/本科/硕士研究生/博士研究生，从档案自动填写"},
            {"name": "毕业学校",                  "fill": "auto", "required_for": [],
             "note": "OCR识别后自动填写"},
            {"name": "专业",                      "fill": "auto", "required_for": [],
             "note": "OCR识别后自动填写"},
            # ── 详情弹窗内更多字段 ──
            {"name": "学习形式",                  "fill": "auto", "required_for": [],
             "note": "全日制/在职/函授，从档案自动填写"},
            {"name": "学制（年）",                "fill": "auto", "required_for": [],
             "note": "如：4"},
            {"name": "学历证书编号",              "fill": "auto", "required_for": [],
             "note": "证书编号，OCR识别后填写"},
            {"name": "学位",                      "fill": "auto", "required_for": [],
             "note": "无/学士/硕士/博士，从档案自动填写"},
            {"name": "学位证书编号",              "fill": "auto", "required_for": [],
             "note": "学位证书编号，OCR识别后填写"},
            {"name": "学位授予单位",              "fill": "auto", "required_for": [],
             "note": "授予学位的学校，OCR识别后填写"},
        ],
    },
    {
        "tab": "3-1",
        "tab_name": "现任专业技术资格、职业资格和技能证书",
        "folder": "03-1_现任专业技术资格",
        "items": [
            # ── 扫描件 ──
            {"name": "职称证书扫描件",            "fill": "file", "required_for": [],
             "note": "OCR识别职称级别/证书编号/资格取得时间/批准机关"},
            {"name": "执业资格证书扫描件（如有）", "fill": "file", "required_for": [],
             "note": "注册建造师/监理工程师等，可放入附_执业资格证目录"},
            # ── 系统自动填写字段（对应网站列表各列）──
            {"name": "数据来源",                  "fill": "auto", "required_for": [],
             "note": "职称证书OCR识别后标记为【数据获取】"},
            {"name": "是否以该资格申报",          "fill": "auto", "required_for": [],
             "note": "当前申报职称的上一级证书填【是】，其余填【否】"},
            {"name": "资格类型",                  "fill": "auto", "required_for": [],
             "note": "职称证书/职业资格证书/技能等级证书，一般填【职称证书】"},
            {"name": "现任专业技术职务",          "fill": "auto", "required_for": [],
             "note": "如：工程技术人员/副高级/高级工程师，OCR识别后自动填写"},
            {"name": "专业",                      "fill": "auto", "required_for": [],
             "note": "OCR识别后自动填写"},
            {"name": "资格取得时间",              "fill": "auto", "required_for": [],
             "note": "格式 YYYY-MM，OCR识别后自动填写"},
            {"name": "证书编号",                  "fill": "auto", "required_for": [],
             "note": "OCR识别后自动填写"},
            {"name": "管理号",                    "fill": "auto", "required_for": [],
             "note": "部分证书有管理号，无则留空"},
            {"name": "现任专业技术职务资格批准机关", "fill": "auto", "required_for": [],
             "note": "如：广西壮族自治区人力资源和社会保障厅，OCR识别后自动填写"},
            {"name": "证书适用范围",              "fill": "auto", "required_for": [],
             "note": "部分证书有适用范围说明，无则留空"},
            {"name": "获取时间",                  "fill": "auto", "required_for": [],
             "note": "证书有效期截止时间，格式 YYYY-MM"},
        ],
    },
    {
        "tab": "3-2",
        "tab_name": "破格/直接申报",
        "folder": "03-2_破格直接申报",
        "items": [
            # ── 第一区块：破格/直接申报 ──
            {"name": "是否申请破格",              "fill": "auto", "required_for": [],
             "note": "从申报人档案填写（是/否），大多数人填【否】"},
            {"name": "破格申报材料（如申请破格）", "fill": "file", "required_for": [],
             "note": "仅申请破格人员需要，放入03-2_破格直接申报目录"},
            # ── 第二区块：职称外语和职称计算机 ──
            {"name": "职称外语计算机要求",        "fill": "auto", "required_for": [],
             "note": "由评委会政策决定，一般为【不作要求】，系统自动填写"},
            {"name": "外语考试合格情况",          "fill": "auto", "required_for": [],
             "note": "由评委会政策决定，一般为【不作要求】，系统自动填写"},
        ],
    },
    {
        "tab": "4",
        "tab_name": "职称外语和职称计算机",
        "folder": "04_外语和计算机",
        "items": [
            # ── 系统自动填写（由评委会政策决定，多数人均为"不作要求"）──
            {"name": "职称外语计算机要求",              "fill": "auto", "required_for": [],
             "note": "由评委会政策决定，默认【不作要求】"},
            {"name": "外语考试合格情况",                "fill": "auto", "required_for": [],
             "note": "默认【不作要求】"},
            {"name": "计算机考试合格情况",              "fill": "auto", "required_for": [],
             "note": "默认【不作要求】"},
            {"name": "职称外语和计算机成绩核查结果",    "fill": "auto", "required_for": [],
             "note": "默认【不作要求】"},
            {"name": "参加外省职称外语、计算机考试的情况说明及查询路径", "fill": "auto", "required_for": [],
             "note": "默认【不作要求】，如有外省考试经历填写说明"},
            # ── 可选扫描件 ──
            {"name": "职称外语和职称计算机免试证明材料", "fill": "file", "required_for": [],
             "note": "有免试资格的才需提供，放入04_外语和计算机目录；无则不需要"},
        ],
    },
    {
        "tab": "5",
        "tab_name": "继续教育学时、学分完成情况",
        "folder": "05_继续教育",
        "items": [
            # ── 扫描件（放入05_继续教育目录）──
            {"name": "继续教育学时证明/培训证书",  "fill": "file", "required_for": [],
             "note": "近年度继续教育凭证，OCR识别年度/学时/机构，每年度放一份"},
            # ── 系统按年度逐行自动填写（对应网站列表各列）──
            {"name": "数据来源",                  "fill": "auto", "required_for": [],
             "note": "固定为【数据获取】，系统自动标记"},
            {"name": "年度",                      "fill": "auto", "required_for": [],
             "note": "如：2024、2025，每行对应一个年度"},
            {"name": "公需必修学时",              "fill": "auto", "required_for": [],
             "note": "从继续教育记录自动填写"},
            {"name": "公需选修学时",              "fill": "auto", "required_for": [],
             "note": "从继续教育记录自动填写"},
            {"name": "行业内数据共享学分",        "fill": "auto", "required_for": [],
             "note": "从继续教育记录自动填写，无则填0"},
            {"name": "行业内数据共享学时",        "fill": "auto", "required_for": [],
             "note": "从继续教育记录自动填写（含住建厅共享数据）"},
            {"name": "专业学时",                  "fill": "auto", "required_for": [],
             "note": "从继续教育记录自动填写"},
            {"name": "总学时（合计）",            "fill": "auto", "required_for": [],
             "note": "系统自动计算各学时之和，需≥90学时/年"},
        ],
    },
    {
        "tab": "6-1",
        "tab_name": "工作简历",
        "folder": "06-1_工作简历",
        "items": [
            {"name": "工作简历（逐行填表）",      "fill": "auto", "required_for": [],
             "note": "网页直接填表，从申报人档案工作经历自动填写，无需扫描件"},
        ],
    },
    {
        "tab": "6-2",
        "tab_name": "个人社保缴纳记录",
        "folder": "06-2_社保记录",
        "items": [
            {"name": "社保缴纳证明/截图",         "fill": "file", "required_for": [],
             "note": "需体现参保单位和起止时间，OCR自动解析时段"},
            {"name": "社保记录（逐行）",          "fill": "auto", "required_for": [],
             "note": "从系统数据自动填写"},
        ],
    },
    {
        "tab": "7-1",
        "tab_name": "专业技术工作经历",
        "folder": "07-1_专业技术工作经历",
        "items": [
            {"name": "业绩证明材料",              "fill": "file", "required_for": [],
             "note": "建设单位出具的专业技术工作经历证明"},
            {"name": "专业技术工作经历（逐行）",  "fill": "auto", "required_for": [],
             "note": "从系统数据自动填写"},
        ],
    },
    {
        "tab": "7-2",
        "tab_name": "学术团体及社会兼职",
        "folder": "07-2_学术团体社会兼职",
        "items": [
            {"name": "学术团体/社会兼职证明（如有）", "fill": "file", "required_for": [],
             "note": "担任学会/协会职务等，可选"},
        ],
    },
    {
        "tab": "8-1",
        "tab_name": "业绩成果",
        "folder": "08-1_业绩成果",
        "items": [
            {"name": "施工合同封面及签章页",      "fill": "file", "required_for": ["高级工程师", "正高级工程师"],
             "note": "高级及以上必须提供"},
            {"name": "竣工验收报告首页",          "fill": "file", "required_for": ["高级工程师", "正高级工程师"],
             "note": "高级及以上必须提供"},
            {"name": "业绩成果（逐行）",          "fill": "auto", "required_for": [],
             "note": "从系统数据自动填写"},
        ],
    },
    {
        "tab": "8-2",
        "tab_name": "获奖情况",
        "folder": "08-2_获奖情况",
        "items": [
            {"name": "获奖证书",                  "fill": "file", "required_for": ["正高级工程师"],
             "note": "正高级必须提供；中高级可选"},
            {"name": "获奖记录（逐行）",          "fill": "auto", "required_for": [],
             "note": "从系统数据自动填写"},
        ],
    },
    {
        "tab": "9",
        "tab_name": "学术成果",
        "folder": "09_学术成果",
        "items": [
            {"name": "论文首页+目录页",           "fill": "file", "required_for": ["高级工程师", "正高级工程师"],
             "note": "高级及以上必须提供"},
            {"name": "期刊收录证明（可选）",      "fill": "file", "required_for": [],
             "note": "核心期刊收录证明"},
            {"name": "学术成果（逐行）",          "fill": "auto", "required_for": [],
             "note": "从系统数据自动填写"},
        ],
    },
    {
        "tab": "10",
        "tab_name": "专业技术工作总结",
        "folder": "10_专业技术工作总结",
        "items": [
            {"name": "专业技术工作总结",          "fill": "ai", "required_for": [],
             "note": "由AI写作模块生成，导出Word后放入10_专业技术工作总结目录，网站粘贴正文"},
        ],
    },
    {
        "tab": "11",
        "tab_name": "其他材料",
        "folder": "11_其他材料",
        "items": [
            {"name": "其他补充材料（如有）",      "fill": "file", "required_for": [],
             "note": "评委会要求的其他材料"},
        ],
    },
    {
        "tab": "附",
        "tab_name": "执业资格证",
        "folder": "附_执业资格证",
        "items": [
            {"name": "注册证书（如有）",          "fill": "file", "required_for": [],
             "note": "注册建造师/监理工程师等，配合3-1填写"},
        ],
    },
]

def _get_materials_for_level(apply_level: str) -> list[dict]:
    """按申报级别过滤材料清单，标注必填/可选。"""
    result = []
    for tab in _TAB_MATERIALS:
        tab_items = []
        for item in tab["items"]:
            required_for = item.get("required_for", [])
            # required_for=[] 表示所有级别必填（fill=auto/ai的视为"系统处理"）
            if item["fill"] in ("auto", "ai"):
                is_required = False   # 系统自动处理，不列为申报人需提交
            elif not required_for:
                is_required = True    # 所有级别必填
            else:
                is_required = apply_level in required_for
            tab_items.append({
                "tab":      tab["tab"],
                "tab_name": tab["tab_name"],
                "folder":   tab["folder"],
                "name":     item["name"],
                "fill":     item["fill"],
                "required": is_required,
                "note":     item.get("note", ""),
            })
        result.extend(tab_items)
    return result

# 向后兼容旧的 _MATERIALS 调用
_MATERIALS = {
    level: _get_materials_for_level(level)
    for level in ["工程师", "高级工程师", "正高级工程师"]
}

@bp.route("/<int:pid>/prescreen")
def prescreen_page(pid):
    project   = get_project(pid)
    if not project:
        flash("项目不存在", "danger")
        return redirect(url_for("projects.list_page"))
    aid       = project.get("applicant_id")
    applicant = get_applicant(aid) if aid else {}
    if not applicant:
        flash("申报人不存在", "danger")
        return redirect(url_for("projects.list_page"))

    from core.prescreen import prescreen, summary as ps_summary
    achievements  = list_achievements(aid)
    awards        = list_awards(aid)
    papers        = list_papers(aid)
    insurances    = list_social_insurance(aid)
    edu_trainings = list_edu_trainings(aid)

    items = prescreen(applicant, project, achievements, awards, papers, insurances, edu_trainings)
    stat  = ps_summary(items)
    return render_template("projects/prescreen.html",
                           project=project, applicant=applicant,
                           items=items, stat=stat)

@bp.route("/<int:pid>/checklist")
def checklist_page(pid):
    project   = get_project(pid)
    if not project:
        flash("项目不存在", "danger")
        return redirect(url_for("projects.list_page"))
    aid       = project.get("applicant_id")
    applicant = get_applicant(aid) if aid else {}
    level     = project.get("apply_level", "")
    materials = _MATERIALS.get(level, _MATERIALS.get("工程师", []))

    # 检查哪些文件已存在于资料目录
    folder = applicant.get("folder_path", "") if applicant else ""
    file_status = _scan_material_files(folder, materials)

    return render_template("projects/checklist.html",
                           project=project, applicant=applicant or {},
                           materials=materials, file_status=file_status)

@bp.route("/<int:pid>/checklist/download")
def checklist_download(pid):
    project   = get_project(pid)
    if not project:
        flash("项目不存在", "danger")
        return redirect(url_for("projects.list_page"))
    aid       = project.get("applicant_id")
    applicant = get_applicant(aid) if aid else {}
    level     = project.get("apply_level", "")
    materials = _MATERIALS.get(level, _MATERIALS.get("工程师", []))
    folder    = applicant.get("folder_path", "") if applicant else ""
    file_status = _scan_material_files(folder, materials)

    from io import BytesIO
    from flask import make_response
    buf = _build_checklist_word(applicant or {}, project, materials, file_status)
    name = (applicant.get("name") or "申报人") if applicant else "申报人"
    resp = make_response(buf.read())
    resp.headers["Content-Type"] = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    resp.headers["Content-Disposition"] = (
        f"attachment; filename=材料清单_{name}.docx")
    return resp


@bp.route("/<int:pid>/fill-preview")
def fill_preview_page(pid):
    project   = get_project(pid)
    if not project:
        flash("项目不存在", "danger")
        return redirect(url_for("projects.list_page"))
    aid       = project.get("applicant_id")
    applicant = get_applicant(aid) if aid else {}

    # Build field mapping: website label → DB value
    a = applicant or {}
    p = dict(project)
    edu_map = {"大专": "专科", "本科": "本科", "硕士研究生": "硕士", "博士研究生": "博士"}
    edu_display = edu_map.get(a.get("education",""), a.get("education",""))

    from datetime import date
    cur_year = date.today().year
    work_age = ""
    if a.get("work_start_year"):
        try: work_age = f"{cur_year - int(a['work_start_year'])} 年"
        except Exception: pass

    fields = [
        {"section": "Tab1·基本信息", "items": [
            {"label": "姓名",           "value": a.get("name",""),          "status": "ok" if a.get("name") else "missing"},
            {"label": "证件号码",       "value": a.get("id_card",""),       "status": "ok" if a.get("id_card") else "missing"},
            {"label": "曾用名",         "value": a.get("former_name",""),   "status": "ok"},
            {"label": "性别",           "value": a.get("gender",""),        "status": "ok" if a.get("gender") else "missing"},
            {"label": "民族",           "value": a.get("ethnicity",""),     "status": "ok" if a.get("ethnicity") else "warn"},
            {"label": "出生年月",       "value": a.get("birth_date",""),    "status": "ok" if a.get("birth_date") else "missing"},
            {"label": "籍贯",           "value": a.get("native_place",""),  "status": "ok" if a.get("native_place") else "warn"},
            {"label": "政治面貌",       "value": a.get("politics",""),      "status": "ok" if a.get("politics") else "warn"},
            {"label": "参加工作时间",   "value": a.get("work_start_date","") or str(a.get("work_start_year","") or ""),
             "status": "ok" if (a.get("work_start_date") or a.get("work_start_year")) else "missing"},
            {"label": "联系电话",       "value": a.get("phone",""),         "status": "ok" if a.get("phone") else "warn"},
            {"label": "电子邮箱",       "value": a.get("email",""),         "status": "ok" if a.get("email") else "warn"},
            {"label": "个人身份性质",   "value": a.get("identity_type",""), "status": "ok" if a.get("identity_type") else "warn"},
            {"label": "拟评职称系列",   "value": p.get("title_series",""),  "status": "ok" if p.get("title_series") else "warn"},
            {"label": "拟评级别",       "value": p.get("apply_level",""),   "status": "ok" if p.get("apply_level") else "missing"},
            {"label": "拟评专业技术资格", "value": p.get("apply_title",""), "status": "ok" if p.get("apply_title") else "warn"},
            {"label": "学科",           "value": p.get("discipline",""),    "status": "ok" if p.get("discipline") else "warn"},
            {"label": "拟评专业",       "value": p.get("specialty",""),     "status": "ok" if p.get("specialty") else "warn"},
            {"label": "专业技术工作年限", "value": work_age,               "status": "ok" if work_age else "missing"},
            {"label": "申报方式",       "value": p.get("apply_method",""),  "status": "ok" if p.get("apply_method") else "warn"},
            {"label": "是否第一次申报", "value": a.get("is_first_apply",""), "status": "ok" if a.get("is_first_apply") else "warn"},
            {"label": "曾申报次数",     "value": str(a.get("apply_count","") or "0"), "status": "ok"},
            {"label": "单位级别",       "value": a.get("unit_level",""),    "status": "ok" if a.get("unit_level") else "warn"},
            {"label": "行政职务",       "value": a.get("admin_position",""), "status": "ok" if a.get("admin_position") else "warn"},
            {"label": "行政职务任命时间", "value": a.get("admin_position_date",""), "status": "ok"},
            {"label": "档案所在地机构名称", "value": a.get("archive_org",""), "status": "ok" if a.get("archive_org") else "warn"},
            {"label": "联系地址",       "value": a.get("address",""),       "status": "ok" if a.get("address") else "warn"},
            {"label": "参加评审会",     "value": p.get("committee",""),     "status": "ok" if p.get("committee") else "warn"},
            {"label": "评审年度",       "value": str(p.get("batch_year","") or ""), "status": "ok" if p.get("batch_year") else "warn"},
        ]},
        {"section": "Tab2·学历情况（列表每行字段）", "items": [
            {"label": "数据来源",     "value": "自填",               "status": "ok"},
            {"label": "毕业时间",     "value": a.get("grad_month","") or (a.get("graduation_year","") + "-" if a.get("graduation_year") else ""),
             "status": "ok" if (a.get("grad_month") or a.get("graduation_year")) else "missing"},
            {"label": "学历项目",     "value": "最高学历",           "status": "ok"},
            {"label": "学历",         "value": edu_display,          "status": "ok" if a.get("education") else "missing"},
            {"label": "毕业学校",     "value": a.get("school",""),   "status": "ok" if a.get("school") else "missing"},
            {"label": "专业",         "value": a.get("major",""),    "status": "ok" if a.get("major") else "missing"},
            {"label": "学习形式",     "value": a.get("study_mode",""), "status": "ok" if a.get("study_mode") else "warn"},
            {"label": "学历证书编号", "value": a.get("edu_cert_no",""), "status": "ok" if a.get("edu_cert_no") else "warn"},
            {"label": "学位",         "value": a.get("degree",""),   "status": "ok" if a.get("degree") else "warn"},
            {"label": "学位证书编号", "value": a.get("degree_cert_no",""), "status": "ok" if a.get("degree_cert_no") else "warn"},
            {"label": "学位授予单位", "value": a.get("degree_school","") or a.get("school",""),
             "status": "ok" if (a.get("degree_school") or a.get("school")) else "warn"},
        ]},
        {"section": "工作单位信息", "items": [
            {"label": "工作单位",   "value": a.get("work_unit",""),       "status": "ok" if a.get("work_unit") else "missing"},
            {"label": "单位类型",   "value": a.get("work_unit_type",""),  "status": "ok" if a.get("work_unit_type") else "warn"},
            {"label": "单位地址",   "value": a.get("work_unit_addr",""),  "status": "ok" if a.get("work_unit_addr") else "warn"},
            {"label": "单位电话",   "value": a.get("work_unit_phone",""), "status": "ok" if a.get("work_unit_phone") else "warn"},
            {"label": "参加工作年份", "value": str(a.get("work_start_year","") or ""),
             "status": "ok" if a.get("work_start_year") else "missing"},
            {"label": "工作年限",   "value": work_age,               "status": "ok" if work_age else "missing"},
            {"label": "现任职务",   "value": a.get("current_position",""),"status": "ok" if a.get("current_position") else "warn"},
            {"label": "现从事专业", "value": a.get("current_specialty",""),"status": "ok" if a.get("current_specialty") else "warn"},
        ]},
        {"section": "Tab3-1·现任专业技术资格、职业资格和技能证书", "items": [
            {"label": "数据来源",       "value": "数据获取",                "status": "ok"},
            {"label": "是否以该资格申报", "value": "是",                   "status": "ok"},
            {"label": "资格类型",       "value": "职称证书",               "status": "ok"},
            {"label": "现任专业技术职务", "value": a.get("title_level",""), "status": "ok" if a.get("title_level") else "missing"},
            {"label": "专业",           "value": a.get("title_specialty",""), "status": "ok" if a.get("title_specialty") else "warn"},
            {"label": "资格取得时间",   "value": a.get("title_month","") or (str(a.get("title_year","")) + "-" if a.get("title_year") else ""),
             "status": "ok" if (a.get("title_month") or a.get("title_year")) else "missing"},
            {"label": "证书编号",       "value": a.get("title_cert_no",""),  "status": "ok" if a.get("title_cert_no") else "warn"},
            {"label": "管理号",         "value": a.get("title_manage_no",""), "status": "ok"},
            {"label": "批准机关",       "value": a.get("title_issuer",""),   "status": "ok" if a.get("title_issuer") else "warn"},
            {"label": "证书适用范围",   "value": a.get("title_scope",""),    "status": "ok"},
            {"label": "获取时间（有效期）", "value": a.get("title_expire",""), "status": "ok"},
        ]},
        {"section": "Tab3-2·破格/直接申报", "items": [
            {"label": "是否申请破格", "value": a.get("apply_exception","否"), "status": "ok"},
        ]},
        {"section": "Tab4·职称外语和职称计算机", "items": [
            {"label": "职称外语计算机要求",   "value": a.get("lang_comp_require","不作要求"),   "status": "ok"},
            {"label": "外语考试合格情况",     "value": a.get("lang_exam_result","不作要求"),    "status": "ok"},
            {"label": "计算机考试合格情况",   "value": a.get("comp_exam_result","不作要求"),    "status": "ok"},
            {"label": "外语和计算机成绩核查结果", "value": a.get("lang_comp_check","不作要求"), "status": "ok"},
            {"label": "参加外省考试情况说明", "value": a.get("lang_comp_other_prov","不作要求"),"status": "ok"},
        ]},
        {"section": "Tab5·继续教育学时、学分完成情况", "items": [
            {"label": "继续教育记录行数",
             "value": f"{len(list_edu_trainings(aid))} 条年度记录" if aid else "—",
             "status": "ok" if aid and list_edu_trainings(aid) else "warn"},
            {"label": "总学时合计",
             "value": str(sum((r.get("total_hours") or
                               (r.get("mandatory_public_hours",0) or 0) +
                               (r.get("elective_public_hours",0) or 0) +
                               (r.get("industry_shared_hours",0) or 0) +
                               (r.get("professional_hours",0) or 0))
                              for r in list_edu_trainings(aid)) if aid else 0) + " 学时",
             "status": "ok"},
            {"label": "每年度≥90学时",
             "value": "请在继续教育记录中逐年核对",
             "status": "warn"},
        ]},
    ]

    missing = sum(1 for sec in fields for it in sec["items"] if it["status"]=="missing")
    warn    = sum(1 for sec in fields for it in sec["items"] if it["status"]=="warn")
    ok      = sum(1 for sec in fields for it in sec["items"] if it["status"]=="ok")

    return render_template("projects/fill_preview.html",
                           project=project, applicant=a,
                           fields=fields, missing=missing, warn=warn, ok=ok)


def _scan_material_files(folder: str, materials: list) -> dict:
    """
    检查资料目录中各Tab子目录是否已有文件。
    返回 {材料名: True/False}
    auto/ai类型的材料直接标记为True（系统处理，无需文件）。
    """
    from core.folder_manager import tab_file_status
    tab_status = tab_file_status(folder) if folder else {}

    result = {}
    for m in materials:
        if m.get("fill") in ("auto", "ai"):
            result[m["name"]] = True   # 系统自动处理
        else:
            folder_name = m.get("folder", "")
            result[m["name"]] = tab_status.get(folder_name, False)
    return result


def _build_checklist_word(applicant: dict, project: dict,
                           materials: list, file_status: dict):
    """生成材料清单 Word 文档，返回 BytesIO。"""
    from io import BytesIO
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from datetime import date

    doc = Document()

    # 页边距
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2)
        sec.left_margin = sec.right_margin = Cm(2.5)

    def _heading(text, size=16, bold=True, center=False):
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        return p

    def _para(text, size=11, bold=False, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)
        return p

    # 标题
    _heading("职称申报材料清单", 16, True, True)
    doc.add_paragraph()

    # 申报人信息表
    info_table = doc.add_table(rows=3, cols=4)
    info_table.style = "Table Grid"
    cells = [
        ("姓    名", applicant.get("name","") or ""),
        ("工作单位", applicant.get("work_unit","") or ""),
        ("申报级别", project.get("apply_level","") or ""),
        ("申报专业", project.get("specialty","") or ""),
        ("现职称",   applicant.get("title_level","") or ""),
        ("截止日期", project.get("deadline","") or ""),
    ]
    for i, (label, val) in enumerate(cells):
        row, col = divmod(i, 2)
        info_table.rows[row].cells[col*2].text     = label
        info_table.rows[row].cells[col*2+1].text   = val
    doc.add_paragraph()

    # 清单标题
    _para(f"请按以下清单准备材料（共 {len(materials)} 项），将文件扫描后放入对应目录：",
          size=11, color=(80,80,80))
    doc.add_paragraph()

    # 材料表格
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, h in enumerate(["序号","材料名称","是否必须","状态"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].font.bold = True

    received = sum(1 for v in file_status.values() if v)
    for idx, m in enumerate(materials, 1):
        row = tbl.add_row().cells
        name   = m["name"]
        found  = file_status.get(name, False)
        row[0].text = str(idx)
        row[1].text = name
        row[2].text = "必须" if m["required"] else "建议"
        row[3].text = "✅ 已收到" if found else "⬜ 待提交"
        if not found and m["required"]:
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.color.rgb = RGBColor(180,0,0)

    doc.add_paragraph()
    _para(f"已收到 {received}/{len(materials)} 项材料", bold=True,
          color=(22,163,74) if received == len(materials) else (180,0,0))
    _para(f"生成时间：{date.today()}", size=10, color=(120,120,120))

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

@bp.route("/checker/scan", methods=["POST"])
def checker_scan():
    import os
    data = request.get_json(silent=True) or {}
    pid = data.get("project_id")
    if not pid:
        return jsonify({"materials": [], "scan": {}, "folder": None})
    project = get_project(pid)
    if not project:
        return jsonify({"materials": [], "scan": {}, "folder": None})

    level = project.get("apply_level", "")
    materials = _MATERIALS.get(level, _MATERIALS.get("工程师", []))
    folder = project.get("folder_path", "")
    scan = {}

    if folder and os.path.isdir(folder):
        try:
            files = os.listdir(folder)
            files_lower = [f.lower() for f in files]
            for m in materials:
                name = m["name"]
                scan[name] = any(name in f or name.replace("复印件","") in f for f in files)
        except Exception:
            pass

    return jsonify({"materials": materials, "scan": scan, "folder": folder})
