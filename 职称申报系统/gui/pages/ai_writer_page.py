"""AI 辅助写作页 — 流式输出 DeepSeek 生成的职称申报文书。"""
from __future__ import annotations
from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QComboBox, QTextEdit, QGroupBox, QFormLayout, QLineEdit,
    QSplitter, QFrame, QMessageBox, QFileDialog, QApplication
)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont
from core.ai_writer import (
    AIWriter, DOCUMENT_TYPES,
    build_work_summary_info, build_masterwork_info, build_achievement_info,
)
from database.models import list_applicants, list_projects
from config import load_config
from core.crypto import decrypt
from gui.widgets import SectionTitle


class StreamWorker(QThread):
    chunk   = pyqtSignal(str)
    done    = pyqtSignal()
    error   = pyqtSignal(str)

    def __init__(self, api_key: str, doc_type: str, info: str, model: str):
        super().__init__()
        self.api_key  = api_key
        self.doc_type = doc_type
        self.info     = info
        self.model    = model

    def run(self):
        try:
            writer = AIWriter(self.api_key, self.model)
            for text in writer.generate_stream(self.doc_type, self.info):
                self.chunk.emit(text)
            self.done.emit()
        except Exception as e:
            self.error.emit(str(e))


class AIWriterPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._projects = []
        self._build_ui()

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(20, 20, 20, 20)
        root.setSpacing(14)
        root.addWidget(SectionTitle("AI 辅助写作"))

        splitter = QSplitter(Qt.Horizontal)

        # ── 左侧：选择 + 上下文 ───────────────────────────────────
        left = QWidget()
        ll = QVBoxLayout(left)
        ll.setContentsMargins(0, 0, 10, 0)
        ll.setSpacing(10)

        # 文书类型
        type_row = QHBoxLayout()
        type_row.addWidget(QLabel("文书类型："))
        self._cmb_type = QComboBox()
        for key, label in DOCUMENT_TYPES.items():
            self._cmb_type.addItem(label, key)
        self._cmb_type.currentIndexChanged.connect(self._on_type_changed)
        type_row.addWidget(self._cmb_type)
        type_row.addStretch()
        ll.addLayout(type_row)

        # 申报人 + 项目选择
        sel_box = QGroupBox("关联数据（自动填充上下文）")
        sel_box.setStyleSheet(
            "QGroupBox{font-weight:bold;color:#1e293b;border:1px solid #e2e8f0;"
            "border-radius:8px;margin-top:6px;padding-top:10px;}"
            "QGroupBox::title{subcontrol-origin:margin;left:12px;}"
        )
        sel_form = QFormLayout(sel_box)
        sel_form.setSpacing(8)
        sel_form.setLabelAlignment(Qt.AlignRight)

        self._cmb_applicant = QComboBox()
        self._cmb_applicant.currentIndexChanged.connect(self._on_applicant_changed)
        self._cmb_project = QComboBox()
        self._cmb_project.currentIndexChanged.connect(self._on_project_changed)

        btn_autofill = QPushButton("自动填充上下文 →")
        btn_autofill.setObjectName("btnOutline")
        btn_autofill.clicked.connect(self._autofill_context)

        sel_form.addRow("申报人：",  self._cmb_applicant)
        sel_form.addRow("申报项目：", self._cmb_project)
        sel_form.addRow("",          btn_autofill)
        ll.addWidget(sel_box)

        # 上下文编辑框
        ll.addWidget(QLabel("写作背景信息（可手动编辑补充）："))
        self._context_edit = QTextEdit()
        self._context_edit.setPlaceholderText(
            "请填写申报人基本情况、工程项目信息、主要业绩等，\n"
            "或点击「自动填充上下文」从数据库导入。\n\n"
            "信息越详细，生成的文书质量越高。"
        )
        ll.addWidget(self._context_edit)

        btn_gen = QPushButton("✨ 开始生成")
        btn_gen.setObjectName("btnPrimary")
        btn_gen.setFixedHeight(40)
        btn_gen.clicked.connect(self._start_generate)
        ll.addWidget(btn_gen)

        self._gen_status = QLabel("")
        self._gen_status.setStyleSheet("color:#64748b; font-size:12px;")
        ll.addWidget(self._gen_status)

        splitter.addWidget(left)

        # ── 右侧：生成结果 ────────────────────────────────────────
        right = QWidget()
        rl = QVBoxLayout(right)
        rl.setContentsMargins(10, 0, 0, 0)
        rl.setSpacing(8)

        result_header = QHBoxLayout()
        result_header.addWidget(QLabel("生成结果："))
        result_header.addStretch()
        btn_copy = QPushButton("复制全文")
        btn_copy.setObjectName("btnOutline")
        btn_copy.clicked.connect(self._copy_result)
        btn_export = QPushButton("导出 .docx")
        btn_export.setObjectName("btnOutline")
        btn_export.clicked.connect(self._export_docx)
        btn_clear = QPushButton("清空")
        btn_clear.setObjectName("btnOutline")
        btn_clear.clicked.connect(lambda: self._result_edit.clear())
        result_header.addWidget(btn_copy)
        result_header.addWidget(btn_export)
        result_header.addWidget(btn_clear)
        rl.addLayout(result_header)

        self._result_edit = QTextEdit()
        self._result_edit.setPlaceholderText("AI 生成的文书将流式显示在此处…")
        f = QFont("Microsoft YaHei"); f.setPointSize(11)
        self._result_edit.setFont(f)
        rl.addWidget(self._result_edit)

        # word count
        self._word_count = QLabel("字数：0")
        self._word_count.setStyleSheet("color:#94a3b8; font-size:11px;")
        self._result_edit.textChanged.connect(
            lambda: self._word_count.setText(
                f"字数：{len(self._result_edit.toPlainText())}"
            )
        )
        rl.addWidget(self._word_count)

        splitter.addWidget(right)
        splitter.setSizes([440, 560])
        root.addWidget(splitter)

    # ── data loading ──────────────────────────────────────────────
    def refresh(self):
        self._cmb_applicant.blockSignals(True)
        self._cmb_applicant.clear()
        self._cmb_applicant.addItem("-- 选择申报人 --", None)
        for a in list_applicants():
            self._cmb_applicant.addItem(a["name"], a["id"])
        self._cmb_applicant.blockSignals(False)

    def _on_applicant_changed(self, _):
        aid = self._cmb_applicant.currentData()
        self._cmb_project.clear()
        self._cmb_project.addItem("-- 选择项目 --", None)
        if not aid:
            self._projects = []
            return
        self._projects = list_projects(applicant_id=aid)
        for p in self._projects:
            label = f"{p.get('apply_level','')} · {p.get('project_name') or p.get('industry','')}"
            self._cmb_project.addItem(label, p["id"])

    def _on_project_changed(self, _):
        pass

    def _on_type_changed(self, _):
        self._context_edit.clear()
        self._result_edit.clear()

    # ── autofill ──────────────────────────────────────────────────
    def _autofill_context(self):
        doc_type = self._cmb_type.currentData()
        aid = self._cmb_applicant.currentData()

        if not aid:
            QMessageBox.warning(self, "提示", "请先选择申报人"); return

        from database.models import get_applicant
        applicant = get_applicant(aid)
        pid = self._cmb_project.currentData()
        project = next((p for p in self._projects if p["id"] == pid), None)

        if doc_type == "work_summary":
            info = build_work_summary_info(applicant, self._projects)
        elif doc_type == "masterwork_desc":
            if not project:
                QMessageBox.warning(self, "提示", "请先选择申报项目"); return
            info = build_masterwork_info(project)
        else:
            if not project:
                QMessageBox.warning(self, "提示", "请先选择申报项目"); return
            info = build_masterwork_info(project)

        self._context_edit.setPlainText(info)

    # ── generate ──────────────────────────────────────────────────
    def _start_generate(self):
        cfg = load_config()
        api_key = decrypt(cfg["ai"].get("deepseek_api_key", ""))
        if not api_key:
            QMessageBox.warning(
                self, "未配置 API Key",
                "请先在「系统设置」中填写 DeepSeek API Key"
            )
            return

        context = self._context_edit.toPlainText().strip()
        if not context:
            QMessageBox.warning(self, "提示", "请填写或自动导入写作背景信息"); return

        doc_type = self._cmb_type.currentData()
        model    = cfg["ai"].get("model", "deepseek-chat")

        self._result_edit.clear()
        self._gen_status.setText("⏳ 生成中…")
        QApplication.processEvents()

        self._worker = StreamWorker(api_key, doc_type, context, model)
        self._worker.chunk.connect(self._on_chunk)
        self._worker.done.connect(self._on_done)
        self._worker.error.connect(self._on_error)
        self._worker.start()

    def _on_chunk(self, text: str):
        cursor = self._result_edit.textCursor()
        cursor.movePosition(cursor.End)
        cursor.insertText(text)
        self._result_edit.setTextCursor(cursor)
        self._result_edit.ensureCursorVisible()

    def _on_done(self):
        self._gen_status.setText("✅ 生成完成")

    def _on_error(self, msg: str):
        self._gen_status.setText("❌ 生成失败")
        QMessageBox.critical(self, "AI 写作错误", msg)

    # ── export ────────────────────────────────────────────────────
    def _copy_result(self):
        text = self._result_edit.toPlainText()
        if text:
            QApplication.clipboard().setText(text)
            self._gen_status.setText("已复制到剪贴板")

    def _export_docx(self):
        text = self._result_edit.toPlainText()
        if not text:
            QMessageBox.information(self, "提示", "暂无内容可导出"); return
        try:
            from docx import Document
        except ImportError:
            QMessageBox.warning(self, "缺少依赖", "请安装 python-docx：\npip install python-docx")
            return

        doc_type = self._cmb_type.currentText()
        path, _ = QFileDialog.getSaveFileName(
            self, "导出文档", f"{doc_type}.docx", "Word 文档 (*.docx)"
        )
        if not path: return

        doc = Document()
        doc.add_heading(doc_type, level=1)
        for para in text.split("\n"):
            doc.add_paragraph(para)
        doc.save(path)
        QMessageBox.information(self, "导出成功", f"已保存至：\n{path}")
